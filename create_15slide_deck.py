import sys
import io
import pptx
from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE

src_path = r'd:\DL\phase2_output\dl_phase2_complete_v3.pptx'
dst_path = r'd:\DL\phase2_output\dl_phase2_review_15slides.pptx'

src_prs = Presentation(src_path)
dst_prs = Presentation()

dst_prs.slide_width = src_prs.slide_width
dst_prs.slide_height = src_prs.slide_height

keep_indices = [
    0,   # 1. Title Slide
    1,   # 2. Domain & Motivation
    4,   # 3. Base Reference Papers
    5,   # 4. Literature Gap
    8,   # 5. Real-World Impact
    13,  # 6. Dataset & Preprocessing Protocol
    14,  # 7. System Architecture
    23,  # 8. Proposed Active Loss & Formula Audit
    24,  # 9. Base Paper Equations Comparison
    15,  # 10. Hyperparameter Tuning Grid
    16,  # 11. Model Comparison Benchmark
    28,  # 12. Final Evaluation Results
    20,  # 13. Adaptation Work & Performance
    29,  # 14. Ablations & Robustness
    9,   # 15. Conclusion & Project Summary
]

print(f"Selecting {len(keep_indices)} slides from total {len(src_prs.slides)} slides...")

blank_layout = dst_prs.slide_layouts[6]

for rank, idx in enumerate(keep_indices, start=1):
    src_slide = src_prs.slides[idx]
    new_slide = dst_prs.slides.add_slide(blank_layout)
    
    for shape in src_slide.shapes:
        if shape.has_text_frame:
            left, top, width, height = shape.left, shape.top, shape.width, shape.height
            txBox = new_slide.shapes.add_textbox(left, top, width, height)
            tf = txBox.text_frame
            tf.word_wrap = shape.text_frame.word_wrap
            
            for p_idx, p in enumerate(shape.text_frame.paragraphs):
                if p_idx == 0:
                    new_p = tf.paragraphs[0]
                else:
                    new_p = tf.add_paragraph()
                
                new_p.text = p.text
                new_p.alignment = p.alignment
                if p.font:
                    if p.font.size: new_p.font.size = p.font.size
                    if p.font.bold: new_p.font.bold = p.font.bold
                    if p.font.name: new_p.font.name = p.font.name
                    try:
                        if p.font.color and p.font.color.type == MSO_COLOR_TYPE.RGB:
                            new_p.font.color.rgb = p.font.color.rgb
                    except Exception:
                        pass
        elif shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            image_bytes = shape.image.blob
            image_stream = io.BytesIO(image_bytes)
            new_slide.shapes.add_picture(image_stream, shape.left, shape.top, shape.width, shape.height)

dst_prs.save(dst_path)
print(f"Successfully generated 15-slide deck at {dst_path}")
