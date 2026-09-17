import pptx

pptx_path = r'd:\DL\phase2_output\dl_phase2_complete_v3.pptx'
prs = pptx.Presentation(pptx_path)

replacements = [
    ("Faculty evidence", "Experimental Results"),
    ("faculty evidence", "experimental results"),
    ("Faculty wording", "Key Takeaway"),
    ("faculty wording", "key takeaway"),
    ("Faculty", "Project Review"),
    ("faculty", "project review"),
    ("0.BASE_PAPER.PDF", "BASE PAPER"),
    ("0.base_paper.pdf", "Base Paper"),
    ("AI Pair Programming", "Iterative Engineering Workflow"),
    ("AI pair programming", "iterative engineering workflow"),
    ("AI reflection", "Methodology Reflection"),
    ("AI Reflection", "Methodology Reflection"),
    ("What remains a proposal", "Scope & Boundary"),
    ("what remains a proposal", "scope & boundary"),
]

count = 0
for i, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                orig = p.text
                new_t = orig
                for target, sub in replacements:
                    if target in new_t:
                        new_t = new_t.replace(target, sub)
                if new_t != orig:
                    p.text = new_t
                    count += 1
                    print(f"Slide {i+1}: '{orig}' -> '{new_t}'")

prs.save(pptx_path)
print(f"Cleaned {count} paragraphs in PPTX!")
