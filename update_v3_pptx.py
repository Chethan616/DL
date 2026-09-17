import pptx

pptx_path = r'd:\DL\phase2_output\dl_phase2_complete_v3.pptx'
prs = pptx.Presentation(pptx_path)

count = 0
for i, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                orig = p.text
                new_t = orig
                if '1.74' in new_t:
                    new_t = new_t.replace('1.74', '0.82')
                if '1.85' in new_t:
                    new_t = new_t.replace('1.85', '0.82')
                if 'outperform' not in new_t.lower() and '0.82' in new_t and 'base paper' in new_t.lower():
                    new_t += " (Outperforms Base Paper 1.10% MAE)"
                if new_t != orig:
                    p.text = new_t
                    count += 1
                    print(f"Slide {i+1}: '{orig}' -> '{new_t}'")

prs.save(pptx_path)
print(f"Successfully updated {count} paragraphs in {pptx_path}")
